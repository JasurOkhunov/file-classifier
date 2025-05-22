from docx import Document
import openpyxl
from io import BytesIO

def extract_text_from_docx(file_bytes: bytes) -> str:
    try:
        doc = Document(BytesIO(file_bytes))
        return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    except Exception as e:
        print(f"[DOCX ERROR] {e}")
        return ""

def extract_text_from_xlsx(file_bytes: bytes) -> str:
    try:
        wb = openpyxl.load_workbook(filename=BytesIO(file_bytes), data_only=True)
        text = []
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                for cell in row:
                    if cell:
                        text.append(str(cell))
        return "\n".join(text)
    except Exception as e:
        print(f"[XLSX ERROR] {e}")
        return ""



